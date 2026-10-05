"""تولید گزارش (HTML / XLSX / PDF) برای تجهیز، مدیریتی و تحلیل NDT."""
from __future__ import annotations

import html
import io
import os
import sqlite3
from datetime import date
from pathlib import Path

import numpy as np

from .. import config
from ..db import one, rows, scalar
from . import corrosion, engine, ndt


def _fmt(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:.2f}"
    return str(v)


# ------------------------------------------------------------- builders
def management_report(conn: sqlite3.Connection) -> dict:
    k = one(conn, "SELECT COUNT(*) total, SUM(status='Active') active FROM equipment WHERE is_deleted=0")
    a = one(conn, "SELECT ROUND(AVG(health_index),2) h, ROUND(AVG(risk_score),1) r, SUM(risk_level='High') hi, SUM(remaining_life_years<5) l5 FROM risk_assessments")
    top = rows(conn, "SELECT e.tag, e.equipment_type, e.location, r.risk_score, r.risk_level, r.remaining_life_years, r.health_index, r.corrosion_rate_mm_per_year "
                     "FROM risk_assessments r JOIN equipment e USING(equipment_id) WHERE e.is_deleted=0 ORDER BY r.risk_score DESC LIMIT 20")
    pl = rows(conn, "SELECT priority, COUNT(*) n, SUM(overdue) o FROM inspection_plans GROUP BY 1 ORDER BY 1")
    de = rows(conn, "SELECT defect_type, COUNT(*) n FROM ndt_records WHERE defect_type!='None' AND review_status!='rejected' GROUP BY 1 ORDER BY n DESC")
    return {"title": "گزارش مدیریتی یکپارچگی دارایی‌ها", "subtitle": f"تاریخ تولید: {date.today().isoformat()}", "sections": [
        {"heading": "خلاصه مدیریتی", "kv": [("تعداد تجهیزات", k["total"]), ("تجهیزات فعال", k["active"]), ("میانگین شاخص سلامت", a["h"]),
                                           ("میانگین امتیاز ریسک", a["r"]), ("تجهیزات با ریسک بالا", a["hi"]), ("تجهیزات با عمر باقیمانده < ۵ سال", a["l5"])]},
        {"heading": "۲۰ تجهیز با بیشترین ریسک", "table": {"columns": ["تگ", "نوع", "واحد", "ریسک", "سطح", "عمر باقیمانده (سال)", "شاخص سلامت", "نرخ خوردگی (mm/y)"],
                                                          "rows": [list(r.values()) for r in top]}},
        {"heading": "برنامه بازرسی بر اساس اولویت", "table": {"columns": ["اولویت", "تعداد", "معوق"], "rows": [list(r.values()) for r in pl]}},
        {"heading": "عیوب شناسایی‌شده", "table": {"columns": ["نوع عیب", "تعداد"], "rows": [list(r.values()) for r in de]}},
        {"heading": "ملاحظات", "text": "خروجی‌های هوش مصنوعی پیشنهادی هستند و تصمیم نهایی بازرسی با کارشناس ذی‌صلاح است."}]}


def equipment_report(conn: sqlite3.Connection, eid: str) -> dict | None:
    e = one(conn, "SELECT * FROM equipment WHERE equipment_id=? AND is_deleted=0", (eid,))
    if not e:
        return None
    r = one(conn, "SELECT * FROM risk_assessments WHERE equipment_id=?", (eid,)) or {}
    p = one(conn, "SELECT * FROM inspection_plans WHERE equipment_id=?", (eid,)) or {}
    hist = engine.load_histories(conn, [eid]).get(eid, [])
    c = corrosion.analyze(hist, e["nominal_thickness_mm"], e["min_required_thickness_mm"], e["install_date"]) if hist else {"ok": False}
    insp = rows(conn, "SELECT i.inspection_date, i.inspection_type, i.method, i.inspector, (SELECT MIN(thickness_mm) FROM thickness_measurements t WHERE t.inspection_id=i.inspection_id) m "
                      "FROM inspections i WHERE equipment_id=? ORDER BY inspection_date DESC LIMIT 12", (eid,))
    nd = rows(conn, "SELECT ndt_id, date, method, defect_type, result, ROUND(confidence,2) c, review_status FROM ndt_records WHERE equipment_id=? AND defect_type!='None' "
                    "ORDER BY date DESC LIMIT 15", (eid,))
    secs = [{"heading": "مشخصات تجهیز", "kv": [("تگ", e["tag"]), ("شناسه", eid), ("نوع", e["equipment_type"]), ("جنس", e["material"]), ("سیال", e["fluid"]),
                                              ("واحد", e["location"]), ("تاریخ نصب", e["install_date"]), ("فشار طراحی (bar)", e["design_pressure_bar"]),
                                              ("دمای طراحی (°C)", e["design_temp_c"]), ("ضخامت اسمی (mm)", e["nominal_thickness_mm"]),
                                              ("حداقل ضخامت مجاز (mm)", e["min_required_thickness_mm"]), ("بحرانیت", e["criticality"]), ("وضعیت", e["status"])]}]
    if r:
        secs.append({"heading": "ارزیابی ریسک و سلامت", "kv": [("امتیاز ریسک", r["risk_score"]), ("سطح ریسک", r["risk_level"]), ("شاخص سلامت", f"{r['health_index']} ({r['health_status']})"),
                                                              ("نرخ خوردگی حاکم (mm/y)", r["corrosion_rate_mm_per_year"]), ("روند", r["trend"]),
                                                              ("عمر باقیمانده (سال)", f"{r['remaining_life_years']} (بازه {r['rl_low']}–{r['rl_high']})"),
                                                              ("احتمال بحرانی ۱/۳/۵ ساله", f"{r['p_critical_1y']:.0%} / {r['p_critical_3y']:.0%} / {r['p_critical_5y']:.0%}")]})
    if c.get("ok"):
        fc = c["forecast"][::4]
        secs.append({"heading": "پیش‌بینی ضخامت (افق ۵ ساله، بازه ۹۰٪)", "table": {"columns": ["تاریخ", "ضخامت (mm)", "حد پایین", "حد بالا"],
                                                                                      "rows": [[f["date"], f["thickness_mm"], f["lower"], f["upper"]] for f in fc]}})
    if p:
        secs.append({"heading": "برنامه بازرسی پیشنهادی", "kv": [("تاریخ پیشنهادی", p["recommended_date"]), ("اولویت", p["priority"]), ("روش", p["recommended_method"]),
                                                                ("استاندارد مبنا", p["standard"]), ("فاصله بازرسی (سال)", p["interval_years"])], "text": p["rationale"]})
    secs.append({"heading": "سوابق بازرسی (آخرین ۱۲ مورد)", "table": {"columns": ["تاریخ", "نوع", "روش", "بازرس", "کمینه ضخامت (mm)"], "rows": [list(x.values()) for x in insp]}})
    if nd:
        secs.append({"heading": "عیوب NDT", "table": {"columns": ["شناسه", "تاریخ", "روش", "عیب", "نتیجه", "اطمینان", "وضعیت بررسی"], "rows": [list(x.values()) for x in nd]}})
    secs.append({"heading": "ملاحظات", "text": "این گزارش به‌صورت خودکار تولید شده و پیش از اقدام اجرایی باید توسط کارشناس ذی‌صلاح تأیید شود."})
    return {"title": f"گزارش یکپارچگی تجهیز {e['tag']}", "subtitle": f"{eid} — تاریخ: {date.today().isoformat()}", "sections": secs}


def ndt_report(conn: sqlite3.Connection, ndt_id: str) -> dict | None:
    n = one(conn, "SELECT n.*, e.tag FROM ndt_records n JOIN equipment e USING(equipment_id) WHERE ndt_id=?", (ndt_id,))
    if not n:
        return None
    secs = [{"heading": "اطلاعات رکورد", "kv": [("شناسه", ndt_id), ("تجهیز", n["tag"]), ("روش", n["method"]), ("تاریخ", n["date"]), ("بازرس", n["inspector"]),
                                              ("نوع عیب", f"{n['defect_type']} ({ndt.FA.get(n['defect_type'], '')})"), ("نتیجه", n["result"]), ("اطمینان", n["confidence"]),
                                              ("اولویت", n["priority_score"]), ("اندازه/عمق/طول (mm)", f"{n['defect_size_mm']} / {n['defect_depth_mm']} / {n['defect_length_mm']}"),
                                              ("وضعیت بررسی", n["review_status"]), ("بررسی‌کننده", n["reviewed_by"] or "—")]}]
    if n["signal_file"] and (config.DATA_DIR / n["signal_file"]).exists():
        a = ndt.analyze_signal(np.load(config.DATA_DIR / n["signal_file"]), with_trace=False)
        secs.append({"heading": "تحلیل خودکار سیگنال", "kv": [("نوع عیب (AI)", f"{a['defect_type']} — {a['defect_type_fa']}"), ("اطمینان", a["confidence"]), ("SNR", a["snr"]),
                                                              ("اولویت", f"{a['priority_score']} ({a['priority_level']})"), ("نتیجه پیشنهادی", a["suggested_result"])]})
        secs.append({"heading": "اندیکاسیون‌ها", "table": {"columns": ["موقعیت نسبی", "دامنه"], "rows": [[i["position"], i["amplitude"]] for i in a["indications"]]}})
    secs.append({"heading": "تأیید کارشناس", "text": n["review_comment"] or "تأیید نهایی توسط کارشناس ذی‌صلاح الزامی است."})
    return {"title": f"گزارش تحلیل NDT — {ndt_id}", "subtitle": f"تاریخ تولید: {date.today().isoformat()}", "sections": secs}


# ------------------------------------------------------------- renderers
CSS = """body{font-family:Tahoma,'Segoe UI',sans-serif;direction:rtl;margin:2rem auto;max-width:960px;color:#1b2430}
h1{color:#0b4f6c;border-bottom:3px solid #0b4f6c;padding-bottom:.4rem}h2{color:#0b4f6c;margin-top:2rem}
table{border-collapse:collapse;width:100%;font-size:.9rem}th{background:#0b4f6c;color:#fff}th,td{border:1px solid #c9d3de;padding:.35rem .6rem;text-align:right}
tr:nth-child(even) td{background:#f3f7fb}dl{display:grid;grid-template-columns:max-content 1fr;gap:.3rem 1.5rem}dt{font-weight:bold;color:#555}dd{margin:0}
.sub{color:#667}@media print{body{margin:0}}"""


def to_html(rep: dict) -> str:
    out = [f"<!doctype html><html lang='fa' dir='rtl'><head><meta charset='utf-8'><title>{html.escape(rep['title'])}</title><style>{CSS}</style></head><body>",
           f"<h1>{html.escape(rep['title'])}</h1><p class='sub'>{html.escape(rep['subtitle'])}</p>"]
    for s in rep["sections"]:
        out.append(f"<h2>{html.escape(s['heading'])}</h2>")
        if s.get("kv"):
            out.append("<dl>" + "".join(f"<dt>{html.escape(str(k))}</dt><dd>{html.escape(_fmt(v))}</dd>" for k, v in s["kv"]) + "</dl>")
        if s.get("table"):
            t = s["table"]
            out.append("<table><tr>" + "".join(f"<th>{html.escape(c)}</th>" for c in t["columns"]) + "</tr>" +
                       "".join("<tr>" + "".join(f"<td>{html.escape(_fmt(c))}</td>" for c in r) + "</tr>" for r in t["rows"]) + "</table>")
        if s.get("text"):
            out.append(f"<p>{html.escape(s['text'])}</p>")
    out.append("</body></html>")
    return "".join(out)


def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش"
    ws.sheet_view.rightToLeft = True
    ws.append([rep["title"]]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([rep["subtitle"]]); ws.append([])
    for s in rep["sections"]:
        ws.append([s["heading"]]); ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, color="0B4F6C")
        for k, v in s.get("kv", []):
            ws.append([str(k), v if isinstance(v, (int, float)) else _fmt(v)])
        if s.get("table"):
            ws.append(s["table"]["columns"])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B4F6C"); c.alignment = Alignment(horizontal="center")
            for r in s["table"]["rows"]:
                ws.append([x for x in r])
        if s.get("text"):
            ws.append([s["text"]])
        ws.append([])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _font_path() -> str | None:
    for p in [os.getenv("REPORT_FONT", ""), "C:/Windows/Fonts/tahoma.ttf", "C:/Windows/Fonts/arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Regular.ttf",
              "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf", "/Library/Fonts/Arial Unicode.ttf"]:
        if p and Path(p).exists():
            return p
    return None


def pdf_available() -> bool:
    try:
        import arabic_reshaper, bidi.algorithm  # noqa
    except ImportError:
        return False
    return _font_path() is not None


def to_pdf(rep: dict) -> bytes:
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(TTFont("FaFont", _font_path()))
    fa = lambda s: get_display(arabic_reshaper.reshape(str(s)))  # noqa: E731
    st = ParagraphStyle("b", fontName="FaFont", fontSize=9, leading=14, alignment=2)
    h1 = ParagraphStyle("h1", parent=st, fontSize=16, leading=22, textColor=colors.HexColor("#0b4f6c"))
    h2 = ParagraphStyle("h2", parent=st, fontSize=12, leading=18, textColor=colors.HexColor("#0b4f6c"), spaceBefore=10)
    cell = ParagraphStyle("c", parent=st, fontSize=8, leading=11, alignment=1)

    def lines(text, width=95):
        import textwrap
        return [Paragraph(fa(l), st) for l in textwrap.wrap(str(text), width)]

    story = [Paragraph(fa(rep["title"]), h1), Paragraph(fa(rep["subtitle"]), st), Spacer(1, 8)]
    for s in rep["sections"]:
        story.append(Paragraph(fa(s["heading"]), h2))
        if s.get("kv"):
            data = [[Paragraph(fa(_fmt(v)), cell), Paragraph(fa(k), cell)] for k, v in s["kv"]]
            t = Table(data, colWidths=[300, 160])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#c9d3de")), ("BACKGROUND", (1, 0), (1, -1), colors.HexColor("#eef3f8"))]))
            story.append(t)
        if s.get("table"):
            tb = s["table"]
            cols = list(reversed(tb["columns"]))
            hc = ParagraphStyle("hc", parent=cell, textColor=colors.white)
            data = [[Paragraph(fa(c), hc) for c in cols]] + [[Paragraph(fa(_fmt(x)), cell) for x in reversed(r)] for r in tb["rows"]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#c9d3de")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b4f6c")),
                                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f7fb")])]))
            story.append(t)
        if s.get("text"):
            story += lines(s["text"])
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4) if any(len(s.get("table", {}).get("columns", [])) > 6 for s in rep["sections"]) else A4,
                      leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24, title=rep["title"]).build(story)
    return buf.getvalue()
